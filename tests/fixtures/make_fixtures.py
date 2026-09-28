"""Gera em tempo de teste os livros curtos de cada formato (sem binarios no repo)."""
from __future__ import annotations

import random
import zipfile
from pathlib import Path

_INGREDIENTS = [
    "farinha de trigo", "açúcar", "manteiga", "ovos", "leite", "fermento", "sal", "azeite",
    "alho", "cebola", "tomate", "arroz", "feijão", "carne moída", "frango", "queijo",
    "creme de leite", "chocolate", "canela", "limão", "batata", "cenoura", "espinafre",
    "camarão", "coco ralado", "fubá", "polvilho", "mandioca", "abóbora", "gengibre",
]
_UNITS = ["xícaras de", "colheres de sopa de", "g de", "ml de", "unidades de", "dentes de", "pitadas de"]
_DISHES = ["BOLO", "TORTA", "SOPA", "FAROFA", "PUDIM", "ESCONDIDINHO", "CREME", "BOLINHO",
           "MOUSSE", "RISOTO", "CALDO", "QUICHE", "ENSOPADO", "PURÊ", "SUFLÊ"]
_VERBS = ["Misture", "Refogue", "Bata", "Cozinhe", "Asse", "Leve ao forno", "Tempere", "Sirva"]


def recipe_book(seed: int, n: int = 6) -> list[dict]:
    rnd = random.Random(seed)
    recipes, used = [], set()
    while len(recipes) < n:
        dish, main = rnd.choice(_DISHES), rnd.choice(_INGREDIENTS)
        title = f"{dish} DE {main.upper()}"
        if title in used:
            continue
        used.add(title)
        ings = rnd.sample(_INGREDIENTS, 6)
        lines = [f"{rnd.randint(1, 500)} {rnd.choice(_UNITS)} {i}" for i in ings]
        steps = " ".join(
            f"{rnd.choice(_VERBS)} {rnd.choice(ings)} com {rnd.choice(ings)} por {rnd.randint(2, 60)} minutos"
            f" na panela, mexendo a massa com cuidado."
            for _ in range(3)
        )
        recipes.append({"title": title, "ingredients": lines, "method": steps})
    return recipes


def book_text(recipes: list[dict]) -> str:
    parts = ["Este livro reúne receitas caseiras testadas na cozinha da família, "
             "com dicas de forno, panela e molho para o dia a dia."]
    for r in recipes:
        parts += [r["title"], "Ingredientes", "\n".join(r["ingredients"]), "Modo de preparo", r["method"]]
    return "\n\n".join(parts) + "\n"


MANUAL_TEXT = "\n\n".join(
    [
        "MANUAL DE CONFIGURAÇÃO DO ROTEADOR",
        "Este documento descreve a configuração de rede do equipamento modelo X-200, incluindo "
        "endereçamento IP, máscara de sub-rede, servidor DHCP e regras de firewall corporativo.",
        "Acesse o painel administrativo pelo navegador, informe o usuário e a senha padrão e "
        "altere as credenciais imediatamente. Em seguida configure a interface WAN com o endereço "
        "fornecido pelo provedor e reinicie o equipamento para aplicar as alterações.",
        "Para diagnosticar falhas de conectividade, verifique o status dos LEDs, os logs do sistema "
        "e a tabela de rotas. Atualize o firmware sempre que houver uma nova versão disponível.",
    ] * 4
)


# ---------- PDF ----------
def _pdf_escape(s: str) -> bytes:
    raw = s.encode("cp1252", errors="replace")
    return raw.replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)")


def write_text_pdf(path: Path, pages: list[list[str]]) -> None:
    """PDF minimo com texto selecionavel. Linha "" = espaco extra (quebra de paragrafo)."""
    objs: list[bytes] = []
    n_pages = len(pages)
    page_ids = [4 + 2 * i for i in range(n_pages)]
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    kids = " ".join(f"{pid} 0 R" for pid in page_ids).encode()
    objs.append(b"<< /Type /Pages /Kids [" + kids + b"] /Count " + str(n_pages).encode() + b" >>")
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")
    for i, lines in enumerate(pages):
        stream = b"BT /F1 11 Tf 14 TL 60 780 Td\n"
        for line in lines:
            stream += (b"(" + _pdf_escape(line) + b") Tj T*\n") if line else b"T*\n"
        stream += b"ET"
        content_id = page_ids[i] + 1
        objs.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 3 0 R >> >>"
            b" /Contents " + str(content_id).encode() + b" 0 R >>"
        )
        objs.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for n, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{n} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(bytes(out))


def _recipe_pdf_pages(recipes: list[dict]) -> list[list[str]]:
    pages = []
    for r in recipes:
        method_lines = [r["method"][i:i + 85] for i in range(0, len(r["method"]), 85)]
        pages.append([r["title"], "", "Ingredientes", "", *r["ingredients"], "", "Modo de preparo", "",
                      *method_lines])
    return pages


def write_image_pdf(path: Path, n_pages: int, text: str | None = None) -> None:
    from PIL import Image, ImageDraw, ImageFont

    images = []
    for i in range(n_pages):
        img = Image.new("RGB", (1240, 1754), "white")
        draw = ImageDraw.Draw(img)
        if text:
            font = ImageFont.load_default(size=36)
            y = 120
            for line in text.splitlines():
                draw.text((100, y), line, fill="black", font=font)
                y += 52
        else:
            draw.rectangle((100, 100 + i * 10, 1100, 400), outline="black", width=4)
        images.append(img)
    images[0].save(path, "PDF", save_all=True, append_images=images[1:], resolution=150)


def write_partial_pdf(path: Path, text_pdf: Path, image_pdf: Path) -> None:
    import pypdfium2 as pdfium

    text_doc, img_doc = pdfium.PdfDocument(str(text_pdf)), pdfium.PdfDocument(str(image_pdf))
    out = pdfium.PdfDocument.new()
    for i in range(len(text_doc)):
        out.import_pages(text_doc, [i])
        if i < len(img_doc):
            out.import_pages(img_doc, [i])
    out.save(str(path))
    for d in (out, text_doc, img_doc):
        d.close()


# ---------- HTML / EPUB ----------
def recipe_html(recipes: list[dict]) -> str:
    body = []
    for r in recipes:
        items = "".join(f"<li>{i}</li>" for i in r["ingredients"])
        body.append(f"<h2>{r['title']}</h2><h3>Ingredientes</h3><ul>{items}</ul>"
                    f"<h3>Modo de preparo</h3><p>{r['method']}</p>")
    return "".join(body)


def write_epub(path: Path, recipes: list[dict], drm: bool = False) -> None:
    half = len(recipes) // 2
    chapters = [recipes[:half], recipes[half:]]
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        zf.writestr("META-INF/container.xml",
                    '<?xml version="1.0"?><container version="1.0" '
                    'xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles>'
                    '<rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>'
                    "</rootfiles></container>")
        items = "".join(f'<item id="c{i}" href="text/c{i}.xhtml" media-type="application/xhtml+xml"/>'
                        for i in range(len(chapters)))
        spine = "".join(f'<itemref idref="c{i}"/>' for i in range(len(chapters)))
        zf.writestr("OEBPS/content.opf",
                    '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0">'
                    f"<metadata/><manifest>{items}</manifest><spine>{spine}</spine></package>")
        for i, chap in enumerate(chapters):
            zf.writestr(f"OEBPS/text/c{i}.xhtml",
                        '<?xml version="1.0" encoding="utf-8"?><html xmlns="http://www.w3.org/1999/xhtml">'
                        f"<body>{recipe_html(chap)}</body></html>")
        if drm:
            zf.writestr("META-INF/encryption.xml",
                        '<encryption xmlns="urn:oasis:names:tc:opendocument:xmlns:container" '
                        'xmlns:enc="http://www.w3.org/2001/04/xmlenc#"><enc:EncryptedData>'
                        '<enc:EncryptionMethod Algorithm="http://www.w3.org/2001/04/xmlenc#aes128-cbc"/>'
                        "</enc:EncryptedData></encryption>")


# ---------- DOCX / ODT / RTF ----------
def write_docx(path: Path, recipes: list[dict]) -> None:
    import docx

    d = docx.Document()
    for n, r in enumerate(recipes):
        d.add_heading(r["title"], level=1)
        d.add_paragraph("Ingredientes")
        if n == 0:
            table = d.add_table(rows=0, cols=2)
            for line in r["ingredients"]:
                qty, _, name = line.partition(" de ")
                row = table.add_row().cells
                row[0].text, row[1].text = f"{qty} de", name
        else:
            for line in r["ingredients"]:
                d.add_paragraph(line)
        d.add_paragraph("Modo de preparo")
        d.add_paragraph(r["method"])
    d.save(str(path))


def write_odt(path: Path, recipes: list[dict]) -> None:
    from odf.opendocument import OpenDocumentText
    from odf.text import H, P

    doc = OpenDocumentText()
    for r in recipes:
        doc.text.addElement(H(outlinelevel=1, text=r["title"]))
        doc.text.addElement(P(text="Ingredientes"))
        for line in r["ingredients"]:
            doc.text.addElement(P(text=line))
        doc.text.addElement(P(text="Modo de preparo"))
        doc.text.addElement(P(text=r["method"]))
    doc.save(str(path))


def _rtf_escape(s: str) -> str:
    out = []
    for ch in s:
        if ch in "\\{}":
            out.append("\\" + ch)
        elif ord(ch) < 128:
            out.append(ch)
        else:
            out.append("\\'%02x" % ch.encode("cp1252", errors="replace")[0])
    return "".join(out)


def write_rtf(path: Path, text: str) -> None:
    body = "".join(_rtf_escape(line) + "\\par\n" for line in text.splitlines())
    path.write_bytes(("{\\rtf1\\ansi\\ansicpg1252{\\fonttbl{\\f0 Helvetica;}}\\f0\n" + body + "}").encode("ascii"))


def build_all(out: Path) -> dict[str, Path]:
    out.mkdir(parents=True, exist_ok=True)
    f: dict[str, Path] = {}

    def p(name: str) -> Path:
        f[name] = out / name
        return f[name]

    txt_book = recipe_book(1)
    p("receitas.txt").write_bytes(book_text(txt_book).encode("latin-1", errors="replace"))
    p("receitas.md").write_text(book_text(txt_book), encoding="utf-8")  # mesmo conteudo do .txt
    write_text_pdf(p("receitas.pdf"), _recipe_pdf_pages(recipe_book(2)))
    write_epub(p("receitas.epub"), recipe_book(3))
    write_epub(p("drm.epub"), recipe_book(3), drm=True)
    write_docx(p("receitas.docx"), recipe_book(4))
    p("protegido.docx").write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 2048)
    write_odt(p("receitas.odt"), recipe_book(5))
    write_rtf(p("receitas.rtf"), book_text(recipe_book(6)))
    p("receitas.html").write_text(
        "<html><head><script>var rastreio = 1;</script></head><body>"
        "<nav>Menu Inicio Contato</nav>" + recipe_html(recipe_book(7)) + "</body></html>",
        encoding="utf-8",
    )
    p("manual.txt").write_text(MANUAL_TEXT, encoding="utf-8")
    write_image_pdf(p("escaneado.pdf"), 3)
    write_partial_pdf(p("parcial.pdf"), f["receitas.pdf"], f["escaneado.pdf"])
    p("quebrado.pdf").write_bytes(f["receitas.pdf"].read_bytes()[:300])
    return f
